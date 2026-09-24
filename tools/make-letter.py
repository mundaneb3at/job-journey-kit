r"""Producer of record for cover letters: reads a content file and writes a matching .docx + .md.

CLI: make-letter.py <target> [--date "Month D, YYYY"] [--out DIR] [--no-docx] [--content-dir DIR]

Reads <content-dir>\<target>.md (default ROOT\letters): header `key: value` lines up to the
first blank line (stem, posting, subject, required, before, dir), body = the paragraphs after
that, split on blank lines. See `templates\letters\_example.md` for the shape. Writes
<out>\<stem>.docx (base_doc margins 0.7/0.8in, Calibri 10.5, space_after 4; name 16pt bold ACCENT
space_after 1; contact 9.5 space_after 10; date 11 space_after 10; subject 11 bold space_after 10;
each body paragraph 11 space_after 10; "Sincerely," 11 space_after 2; name 11 bold) and
<out>\<stem>.md, the same mirror shape, pointing at this file and its gate (`verify-letter.py`).

The posting is only resolved to fail fast (exit 2) when it is missing; its text is not read here -
the `quotes` probe in verify-letter.py reads it against the built letter body.
"""
import argparse
import datetime
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parents[1]          # kit-user folder (this file lives in ROOT\tools)

# ---------------------------------------------------------------- config (edit for yourself)
NAME = "<Your Display Name>"
CONTACT = "<Your City> | <your email> | <phone> | <github.com/yourhandle>"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
BODY = RGBColor(0x22, 0x22, 0x22)

HEADER_RX = re.compile(r"^([a-z]+):\s*(.*)$")


def base_doc():
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.8)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = BODY
    st.paragraph_format.space_after = Pt(4)
    return d


def para(d, text, size=10.5, bold=False, color=BODY, space_after=4):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def default_date():
    d = datetime.date.today()
    return "%s %d, %d" % (d.strftime("%B"), d.day, d.year)          # no zero-pad, e.g. "September 23, 2026"


def read_content(path):
    """<target>.md -> dict(stem, posting, subject, required[list], before(str|None), body[list])."""
    raw = path.read_text(encoding="utf-8")
    head, rest = re.split(r"\n\s*\n", raw, maxsplit=1)
    cfg = {}
    for ln in head.splitlines():
        m = HEADER_RX.match(ln.strip())
        if m:
            cfg[m.group(1)] = m.group(2).strip()
    cfg["required"] = [r.strip() for r in cfg.get("required", "").split("|") if r.strip()]
    before = cfg.get("before", "none")
    cfg["before"] = None if before.lower() == "none" else before
    cfg["body"] = [" ".join(block.split()) for block in re.split(r"\n\s*\n", rest.strip()) if block.strip()]
    return cfg


def resolve_posting(key, content_dir):
    """-> Path of the posting file, or exit 2 if neither candidate exists."""
    candidates = [content_dir / "postings" / (key + ext) for ext in (".paste.md", ".api.txt")]
    for c in candidates:
        if c.exists():
            return c
    print("no posting file for %r - checked:\n  " % key + "\n  ".join(str(c) for c in candidates))
    sys.exit(2)


def write_docx(path, letter_date, subject, body):
    d = base_doc()
    para(d, NAME, size=16, bold=True, color=ACCENT, space_after=1)
    para(d, CONTACT, size=9.5, space_after=10)
    para(d, letter_date, size=11, space_after=10)
    para(d, subject, size=11, bold=True, space_after=10)
    for t in body:
        para(d, t, size=11, space_after=10)
    para(d, "Sincerely,", size=11, space_after=2)
    para(d, NAME, size=11, bold=True)
    d.save(str(path))
    print("WROTE", path)


def write_mirror(path, letter_date, subject, body):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(
            "# %s — Cover Letter: %s\n\n" % (NAME, subject)
            + "> Generated mirror of `tools\\make-letter.py` (producer of record). Do not edit; "
              "edit the content file and rerun. Upload the .docx. Gate: `tools\\verify-letter.py`.\n\n"
            + "**Date:** %s\n\n**Subject:** %s\n\n" % (letter_date, subject)
            + "\n\n".join(body) + "\n\nSincerely,\n%s\n" % NAME
        )
    print("WROTE", path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("target", help="content file stem under --content-dir, e.g. 'example'")
    ap.add_argument("--date", default=None, help='"Month D, YYYY", default today')
    ap.add_argument("--out", default=None, help="output dir, default the kit-user root")
    ap.add_argument("--no-docx", action="store_true", help="write only the .md mirror")
    ap.add_argument("--content-dir", default=None, help="where <target>.md lives, default ROOT\\letters")
    a = ap.parse_args()

    content_dir = Path(a.content_dir) if a.content_dir else (ROOT / "letters")
    letter_date = a.date or default_date()

    src = content_dir / (a.target + ".md")
    if not src.exists():
        print("no content file: %s" % src)
        sys.exit(2)
    cfg = read_content(src)
    # default output = the `dir:` header (relative to the kit-user root) or the root itself
    out_dir = Path(a.out) if a.out else (ROOT / cfg["dir"] if cfg.get("dir") else ROOT)
    out_dir.mkdir(parents=True, exist_ok=True)
    resolve_posting(cfg["posting"], content_dir)   # fail fast (exit 2) if the posting is missing

    if not a.no_docx:
        write_docx(out_dir / (cfg["stem"] + ".docx"), letter_date, cfg["subject"], cfg["body"])
    write_mirror(out_dir / (cfg["stem"] + ".md"), letter_date, cfg["subject"], cfg["body"])


if __name__ == "__main__":
    main()
