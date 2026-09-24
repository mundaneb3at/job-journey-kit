r"""Build Resume-<version>-<date>.docx (page 1 technical + page 2 optional tail) from a content file.

Layout (measurements, python-docx helpers, build()/check_widths()/render()) lives entirely in this
file; content (every string the resume shows) comes from `--content` (default `resume-content.py`
next to this repo's `tools\` folder) via importlib from a file path, so a new version is
`--version vN --date ... --content resume-content.py` with no code copy. Start from
`templates\resume-content.example.py` for your own content file.

LOOK = a common university career-advisor template shape: Calibri 10pt black, single spacing;
name 14pt bold with a rule under it; section headings 10pt bold underlined with a blank line
before; bullets with a 284-twip (0.2") hanging indent; skills as bullets with a bold label; dates
on a right tab at 10206 twips; margins 680/851 twips. Edit the constants below if your own
template's measurements differ.

  python tools\make-resume.py --version v1 --date 2026-01-15
      -> writes out\Resume-v1-2026-01-15.docx
  python tools\make-resume.py --version v1 --date 2026-01-15 --render
      -> also Word COM -> PDF -> PNG per page (needs Word + pywin32 + pymupdf)
  Refuses (exit 2) to overwrite an existing output docx unless --force.
"""
import argparse
import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

HERE = Path(__file__).resolve().parent          # <kit-user folder>\tools
ROOT = HERE.parent                               # kit-user folder

BLACK = RGBColor(0, 0, 0)
BASE = 10                     # template body size (w:sz 20)
MARGIN_TB, MARGIN_LR = 680, 851   # twips, from the template's sectPr
RIGHT_TAB = 10206             # twips, the template's right tab for dates
BULLET_IND = 284              # twips, the template's hanging indent
PAGE_W = 12240
TEXT_W_PT = (PAGE_W - 2 * MARGIN_LR) / 20      # 526.9 pt
BULLET_W_PT = TEXT_W_PT - BULLET_IND / 20      # 512.7 pt

# ---------------------------------------------------------------- content loading

def load_content(path):
    """Import a content file (default resume-content.py) by path and return the module."""
    spec = importlib.util.spec_from_file_location("resume_content", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

# ---------------------------------------------------------------- helpers

def base_doc():
    d = Document()
    for s in d.sections:
        s.page_width = Twips(PAGE_W); s.page_height = Twips(15840)
        s.top_margin = s.bottom_margin = Twips(MARGIN_TB)
        s.left_margin = s.right_margin = Twips(MARGIN_LR)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(BASE)
    st.font.color.rgb = BLACK
    pf = st.paragraph_format
    pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
    return d


def run(p, text, size=BASE, bold=False, italic=False, underline=False):
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.underline = underline
    r.font.color.rgb = BLACK
    return r


def para(d, text="", size=BASE, bold=False, italic=False, before=0, after=0, style=None, align=None):
    p = d.add_paragraph(style=style)
    if text:
        run(p, text, size, bold, italic)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p


def bottom_border(p, sz=12):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "1"); b.set(qn("w:color"), "auto")
    pbdr.append(b); pPr.append(pbdr)


def heading(d, text):
    # blank 10pt line before (space_before stands in for the empty paragraph)
    p = para(d, before=10)
    run(p, text, bold=True, underline=True)
    return p


def right_tab(p):
    p.paragraph_format.tab_stops.add_tab_stop(Twips(RIGHT_TAB), WD_TAB_ALIGNMENT.RIGHT)


def bullet(d, text, label=None):
    p = d.add_paragraph(style="List Bullet")
    if label:
        run(p, label, bold=True)
    run(p, text)
    pf = p.paragraph_format
    pf.left_indent = Twips(BULLET_IND); pf.first_line_indent = Twips(-BULLET_IND)
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    return p


def hyperlink(p, url, text, size=BASE):
    """A real clickable link (w:hyperlink + external relationship), shown as plain black text.
    python-docx has no public API for this."""
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                            is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size * 2)); rpr.append(sz)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "000000"); rpr.append(col)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(rpr); r.append(t); h.append(r); p._p.append(h)
    return h


def entry(d, title, rest="", date=None, before=2, link=None):
    p = para(d, before=before); right_tab(p)
    run(p, title, bold=True)
    if rest:
        run(p, rest)
    if link:
        hyperlink(p, "https://" + link, link)
    if date:
        run(p, "\t" + date)
    return p


# ---------------------------------------------------------------- width check (one line per bullet)

def text_width_pt(text, size=BASE, bold=False):
    """Rendered width in points using the installed Calibri; None if PIL or the font is missing."""
    try:
        from PIL import ImageFont
        f = ImageFont.truetype(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf", 100)
        return f.getlength(text) / 100 * size
    except Exception:
        return None


def check_widths(c):
    """Print every bullet that would wrap at 10pt. Exit code stays 0; this is a report."""
    long = []
    for h in c.HIGHLIGHTS:
        w = text_width_pt(h)
        if w and w > BULLET_W_PT: long.append(("highlight", round(w), h[:70]))
    for label, rest in c.SKILLS:
        w = (text_width_pt(label, bold=True) or 0) + (text_width_pt(rest) or 0)
        if w > 2 * BULLET_W_PT: long.append(("skill>2 lines", round(w), label))
    # every non-skill bullet on BOTH pages should be one rendered line
    rows = [(t, b) for t, _, _, _, bs in c.PROJECTS for b in bs] + [(t, b) for t, _, _, b in c.EXTRA_PROJECTS]
    rows += [(t, d) for t, _, _, d in c.JOBS if d] + [("extracurricular", c.EXTRACURRICULAR)]
    for title, b in rows:
        w = text_width_pt(b)
        if w and w > BULLET_W_PT: long.append((title, round(w), b[:70]))
    for row in long:
        print("WRAPS  %-16s %4s pt > %d : %s" % (row[0], row[1], BULLET_W_PT, row[2]))
    print("width check: %d bullet(s) wider than one line" % len(long))
    return long


# ---------------------------------------------------------------- build

def build(c, out):
    d = base_doc()

    # header: name 14pt bold with a rule; contact line; no summary/title line
    p = para(d); run(p, c.NAME, 14, bold=True); bottom_border(p)
    p = para(d, size=11); run(p, c.CONTACT_LINE, 11)
    # Links line. Separators are bare runs (no w:b / w:u tags): run() writes <w:b w:val="0"/> and
    # <w:u w:val="none"/>, which verify-template.py reads as bold + underlined, and a line with
    # several such tokens would then wrongly count as a section heading.
    p = para(d, size=11)
    for i, link in enumerate(c.LINKS):
        if i:
            p.add_run(" | ").font.size = Pt(11)
        hyperlink(p, "https://" + link, link, size=11)
    # role + term on one short contact-block line (<= 12 words: verify-template.py reads a longer
    # non-bullet line before the first section as a summary paragraph, which this format forbids).
    para(d, c.TERM_LINE, size=11)

    heading(d, "Highlights")
    for h in c.HIGHLIGHTS:
        bullet(d, h)

    heading(d, "Education")
    entry(d, c.EDUCATION[0], c.EDUCATION[1], c.EDUCATION[2], before=0)

    heading(d, "Technical Skills")
    for label, rest in c.SKILLS:
        bullet(d, rest, label=label)

    heading(d, "Technical Projects")
    for title, sub, link, date, bullets in c.PROJECTS:
        entry(d, title, " - " + sub + " \u00b7 ", date, link=link)
        for b in bullets:
            bullet(d, b)

    # ---- page 2: an optional tail (header: name + contact on a right tab)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = para(d); right_tab(p); run(p, c.NAME, 14, bold=True)
    run(p, c.P2_HEADER_CONTACT); bottom_border(p)

    # heading must not contain "technical projects" - verify-template.py ranks that substring and
    # would misfile this section as a duplicate of the page-1 one.
    heading(d, "Additional Projects")
    for title, link, date, b in c.EXTRA_PROJECTS:
        entry(d, title, " \u00b7 ", date, link=link)
        bullet(d, b)

    heading(d, "Additional Experience")   # non-technical work experience, e.g. general jobs
    for title, org, when, desc in c.JOBS:
        entry(d, title, org, when)
        if desc:
            bullet(d, desc)

    heading(d, "Extracurricular Activities")
    bullet(d, c.EXTRACURRICULAR)

    d.save(out)
    print("WROTE", out)


def render(out):
    """Word COM -> PDF -> one PNG per page, next to the docx. Prints the page count Word reports."""
    import pythoncom  # noqa: F401  (pywin32)
    import win32com.client
    import fitz
    pdf = out.with_suffix(".pdf")
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(out), ReadOnly=True)
        doc.ExportAsFixedFormat(str(pdf), 17)  # wdExportFormatPDF
        pages = doc.ComputeStatistics(2)       # wdStatisticPages
        doc.Close(False)
    finally:
        word.Quit()
    print("PDF", pdf, "pages(Word)=", pages)
    with fitz.open(pdf) as f:
        for i, page in enumerate(f, 1):
            png = out.with_name(f"render-page{i}.png")
            page.get_pixmap(dpi=110).save(png)
            print("PNG", png)
        print("pages(pdf)=", f.page_count)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--version", required=True, help="e.g. v1")
    ap.add_argument("--date", required=True, help="YYYY-MM-DD")
    ap.add_argument("--out", help="output dir (default out\\)")
    ap.add_argument("--content", help="content .py file (default resume-content.py at the kit-user root)")
    ap.add_argument("--render", action="store_true", help="Word COM -> PDF -> PNG after writing")
    ap.add_argument("--force", action="store_true", help="overwrite an existing output docx")
    args = ap.parse_args()

    content_path = Path(args.content).resolve() if args.content else ROOT / "resume-content.py"
    out_dir = Path(args.out).resolve() if args.out else ROOT / "out"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"Resume-{args.version}-{args.date}.docx"

    if out.exists() and not args.force:
        print("REFUSING to overwrite existing file (pass --force):", out)
        sys.exit(2)

    c = load_content(content_path)
    check_widths(c)
    build(c, out)
    if args.render:
        render(out)


if __name__ == "__main__":
    main()
